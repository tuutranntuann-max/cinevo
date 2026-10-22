from io import BytesIO
from datetime import timedelta

from openpyxl import Workbook

from app.extensions import db
from app.models import Booking, BookingSeat, Cinema, Movie, Room, Seat, Showtime, User
from app.timeutil import local_now


def test_province_catalog_lists_all_current_provinces_by_region(client):
    response = client.get("/api/provinces")
    assert response.status_code == 200
    provinces = response.get_json()
    assert len(provinces) == 34
    assert len({province["name"] for province in provinces}) == 34
    assert {region for region in (province["region"] for province in provinces)} == {
        "Miền Bắc", "Miền Trung", "Miền Nam",
    }
    assert sum(province["region"] == "Miền Bắc" for province in provinces) == 15
    assert sum(province["region"] == "Miền Trung" for province in provinces) == 11
    assert sum(province["region"] == "Miền Nam" for province in provinces) == 8


def test_admin_imports_cinema_rooms_and_seats_with_duplicate_skip(client, auth, app):
    csv_data = (
        "Tên rạp,Miền,Tỉnh/thành phố,Địa chỉ,SĐT,Tên phòng,Loại phòng,Số hàng,Số cột\n"
        "Cinevo Cầu Giấy,Miền Bắc,Hà Nội,Tầng 5,02438000001,Phòng mới,2D,2,3\n"
        "Cinevo Đà Nẵng,Miền Trung,Đà Nẵng,123 Hải Châu,02361234567,Phòng 01,2D,2,3\n"
        "Cinevo Đà Nẵng,Miền Trung,Đà Nẵng,123 Hải Châu,02361234567,IMAX,IMAX,3,4\n"
    ).encode("utf-8-sig")
    assert client.post("/api/admin/cinemas/import", data={
        "file": (BytesIO(csv_data), "cinemas.csv"),
    }).status_code == 401
    assert client.post("/api/admin/cinemas/import", headers=auth("manager"), data={
        "file": (BytesIO(csv_data), "cinemas.csv"),
    }).status_code == 403

    response = client.post("/api/admin/cinemas/import", headers=auth("admin"), data={
        "file": (BytesIO(csv_data), "cinemas.csv"),
    })
    assert response.status_code == 200
    assert response.get_json()["imported_count"] == 1
    assert response.get_json()["skipped_count"] == 1

    with app.app_context():
        cinema = Cinema.query.filter_by(name="Cinevo Đà Nẵng").one()
        assert cinema.city == "Thành phố Đà Nẵng"
        rooms = Room.query.filter_by(cinema_id=cinema.id).order_by(Room.name).all()
        assert len(rooms) == 2
        assert {
            room.name: Seat.query.filter_by(room_id=room.id).count()
            for room in rooms
        } == {"Phòng 01": 6, "IMAX": 12}
        assert Cinema.query.filter_by(name="Cinevo Cầu Giấy").count() == 1
        assert Room.query.filter_by(name="Phòng mới").count() == 0

    public_cinemas = client.get("/api/cinemas").get_json()
    imported = next(cinema for cinema in public_cinemas if cinema["name"] == "Cinevo Đà Nẵng")
    assert imported["region"] == "Miền Trung"
    assert imported["city"] == "Thành phố Đà Nẵng"
    legacy_cinema = next(cinema for cinema in public_cinemas if cinema["name"] == "Cinevo Cầu Giấy")
    assert legacy_cinema["province"] == "Thành phố Hà Nội"
    assert legacy_cinema["region"] == "Miền Bắc"


def test_admin_can_add_cinema_and_default_room_manually(client, auth, app):
    response = client.post("/api/admin/cinemas", headers=auth("admin"), json={
        "name": "Cinevo Lào Cai",
        "city": "Lào Cai",
        "address": "10 Hoàng Liên",
        "room_name": "Phòng 01",
        "room_type": "3D",
        "total_rows": 3,
        "total_cols": 4,
    })
    assert response.status_code == 201
    assert response.get_json()["city"] == "Lào Cai"
    assert response.get_json()["region"] == "Miền Bắc"
    assert response.get_json()["rooms"][0]["room_type"] == "3D"
    with app.app_context():
        cinema = Cinema.query.filter_by(name="Cinevo Lào Cai").one()
        assert Seat.query.join(Room).filter(Room.cinema_id == cinema.id).count() == 12

    duplicate = client.post("/api/admin/cinemas", headers=auth("admin"), json={
        "name": "Cinevo Lào Cai",
        "city": "Lào Cai",
        "address": "Địa chỉ khác",
        "room_name": "Phòng 02",
    })
    assert duplicate.status_code == 409


def test_admin_can_edit_cinema_rooms_and_seats_safely(client, auth, app):
    created = client.post("/api/admin/cinemas", headers=auth("admin"), json={
        "name": "Cinevo Sửa Rạp",
        "city": "Hà Nội",
        "address": "Địa chỉ cũ",
        "room_name": "Phòng 01",
        "total_rows": 3,
        "total_cols": 4,
    })
    assert created.status_code == 201
    cinema_id = created.get_json()["id"]
    room_id = created.get_json()["rooms"][0]["id"]

    updated_cinema = client.put(f"/api/admin/cinemas/{cinema_id}", headers=auth("admin"), json={
        "chain": "Cinevo",
        "name": "Cinevo Sửa Rạp Mới",
        "city": "Hà Nội",
        "address": "",
        "phone": "0900000000",
    })
    assert updated_cinema.status_code == 200
    assert updated_cinema.get_json()["address"] is None
    assert updated_cinema.get_json()["name"] == "Cinevo Sửa Rạp Mới"

    added_room = client.post(f"/api/admin/cinemas/{cinema_id}/rooms", headers=auth("admin"), json={
        "name": "Phòng 02",
        "room_type": "3D",
        "total_rows": 2,
        "total_cols": 3,
    })
    assert added_room.status_code == 201
    assert added_room.get_json()["seats_count"] == 6

    expanded = client.put(f"/api/admin/rooms/{room_id}", headers=auth("admin"), json={
        "name": "Phòng 01 Mới",
        "room_type": "IMAX",
        "total_rows": 4,
        "total_cols": 5,
    })
    assert expanded.status_code == 200
    assert expanded.get_json()["seats_count"] == 20

    with app.app_context():
        room = db.session.get(Room, room_id)
        movie = Movie.query.first()
        user = User.query.filter_by(email="user@cinevo.vn").one()
        showtime = Showtime(
            movie_id=movie.id,
            room_id=room.id,
            start_time=local_now() + timedelta(days=1),
            base_price=90000,
        )
        db.session.add(showtime)
        db.session.flush()
        booked_seat = Seat.query.filter_by(room_id=room.id, row_label="D", number=5).one()
        booking = Booking(
            code="CVTEST0001",
            user_id=user.id,
            showtime_id=showtime.id,
            seat_amount=90000,
            total_amount=90000,
            status="CANCELLED",
        )
        db.session.add(booking)
        db.session.flush()
        db.session.add(BookingSeat(
            booking_id=booking.id,
            showtime_id=showtime.id,
            seat_id=booked_seat.id,
            price=90000,
        ))
        db.session.commit()

    unsafe_shrink = client.put(f"/api/admin/rooms/{room_id}", headers=auth("admin"), json={
        "name": "Phòng 01 Mới",
        "room_type": "IMAX",
        "total_rows": 3,
        "total_cols": 4,
    })
    assert unsafe_shrink.status_code == 409
    assert "đã có vé đặt" in unsafe_shrink.get_json()["message"]

    protected_delete = client.delete(f"/api/admin/rooms/{room_id}", headers=auth("admin"))
    assert protected_delete.status_code == 409
    deleted_empty_room = client.delete(f"/api/admin/rooms/{added_room.get_json()['id']}", headers=auth("admin"))
    assert deleted_empty_room.status_code == 200
    with app.app_context():
        assert Seat.query.filter_by(room_id=added_room.get_json()["id"]).count() == 0


def test_cinema_import_rejects_invalid_room_without_partial_write(client, auth, app):
    invalid_csv = (
        "Tên rạp,Miền,Tỉnh/thành phố,Địa chỉ,Tên phòng,Loại phòng,Số hàng,Số cột\n"
        "Should Not Persist,Miền Nam,Đà Nẵng,Địa chỉ,Phòng 01,2D,8,10\n"
        "Another Invalid,Miền Trung,Đà Nẵng,Địa chỉ,Phòng 01,4D,8,10\n"
    ).encode("utf-8")
    with app.app_context():
        original_count = Cinema.query.count()
    response = client.post("/api/admin/cinemas/import", headers=auth("admin"), data={
        "file": (BytesIO(invalid_csv), "bad.csv"),
    })
    assert response.status_code == 400
    with app.app_context():
        assert Cinema.query.count() == original_count

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Tên rạp", "Tỉnh/thành phố", "Địa chỉ", "Tên phòng", "Loại phòng", "Số hàng", "Số cột"])
    sheet.append(["Cinevo Huế", "Huế", "1 Lê Lợi", "Phòng 01", "2D", 2, 5])
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    imported = client.post("/api/admin/cinemas/import", headers=auth("admin"), data={
        "file": (BytesIO(content.getvalue()), "cinemas.xlsx"),
    })
    assert imported.status_code == 200
    with app.app_context():
        cinema = Cinema.query.filter_by(name="Cinevo Huế").one()
        assert cinema.city == "Thành phố Huế"
        assert Room.query.filter_by(cinema_id=cinema.id).count() == 1


def test_cinema_import_accepts_workbook_headers_and_optional_address(client, auth, app):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([
        "STT", "Miền", "Tỉnh/thành", "Hệ thống rạp", "Tên cụm rạp", "Địa chỉ",
        "Tình trạng địa chỉ", "Ghi chú", "Nguồn dữ liệu",
    ])
    sheet.append([
        1, "Miền Bắc", "Bắc Giang", "Beta Cineplex", "Beta Cineplex Bắc Giang",
        None, "Chưa tìm thấy địa chỉ", None, "Nguồn kiểm thử",
    ])
    content = BytesIO()
    workbook.save(content)
    workbook.close()

    response = client.post("/api/admin/cinemas/import", headers=auth("admin"), data={
        "file": (BytesIO(content.getvalue()), "rap-bo-sung.xlsx"),
    })
    assert response.status_code == 200
    assert response.get_json()["imported_count"] == 1
    assert response.get_json()["missing_address_count"] == 1
    with app.app_context():
        cinema = Cinema.query.filter_by(name="Beta Cineplex Bắc Giang").one()
        assert cinema.chain == "Beta Cineplex"
        assert cinema.city == "Bắc Ninh"
        assert cinema.address is None
        assert Room.query.filter_by(cinema_id=cinema.id).one().name == "Phòng 01"


def test_cinema_import_uses_current_province_region_instead_of_stale_excel_region(client, auth):
    csv_data = (
        "Miền,Tỉnh/thành,Hệ thống rạp,Tên cụm rạp\n"
        "Miền Bắc,Thanh Hóa,Cinema Test,Cinema Test Thanh Hóa\n"
    ).encode("utf-8-sig")
    response = client.post("/api/admin/cinemas/import", headers=auth("admin"), data={
        "file": (BytesIO(csv_data), "cinema.csv"),
    })
    assert response.status_code == 200
    assert response.get_json()["corrected_region_count"] == 1
    cinema = next(
        row for row in client.get("/api/admin/cinemas", headers=auth("admin")).get_json()
        if row["name"] == "Cinema Test Thanh Hóa"
    )
    assert cinema["city"] == "Thanh Hóa"
    assert cinema["region"] == "Miền Trung"


def test_admin_imports_users_cinema_excel_headers_and_region_wrapped_legacy_locations(client, auth, app):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Hệ thống rạp", "Tên cụm rạp", "Khu vực / Tỉnh thành", "Địa chỉ / Chi tiết"])
    sheet.append([
        "CGV Cinemas", "CGV Vincom Bà Triệu", "Miền Bắc (Hà Nội)",
        "Tầng 6, TTTM Vincom Center, 191 Bà Triệu, Hai Bà Trưng",
    ])
    sheet.append([
        "CGV Cinemas", "CGV Vincom Đà Nẵng", "Miền Trung (Đà Nẵng)",
        "Tầng 4, TTTM Vincom Plaza Đà Nẵng, 91 Ngô Quyền",
    ])
    sheet.append([
        "CGV Cinemas", "CGV Kim Cương Quy Nhơn", "Miền Trung (Bình Định)",
        "Tòa nhà An Phú Thịnh, 1 Tăng Bạt Hổ, Quy Nhơn",
    ])
    sheet.append([
        "CGV Cinemas", "CGV Aeon Bình Dương", "Miền Nam (Bình Dương)",
        "TTTM Aeon Mall Bình Dương, Thuận An",
    ])
    file_content = BytesIO()
    workbook.save(file_content)
    workbook.close()

    response = client.post("/api/admin/cinemas/import", headers=auth("admin"), data={
        "file": (BytesIO(file_content.getvalue()), "DanhSachRapChieu.xlsx"),
    })
    assert response.status_code == 200
    assert response.get_json()["imported_count"] == 4
    assert response.get_json()["skipped_count"] == 0

    with app.app_context():
        expected = {
            "CGV Vincom Bà Triệu": "Thành phố Hà Nội",
            "CGV Vincom Đà Nẵng": "Thành phố Đà Nẵng",
            "CGV Kim Cương Quy Nhơn": "Gia Lai",
            "CGV Aeon Bình Dương": "Thành phố Hồ Chí Minh",
        }
        for cinema_name, province in expected.items():
            cinema = Cinema.query.filter_by(name=cinema_name).one()
            assert cinema.city == province
            assert cinema.chain == "CGV Cinemas"
            room = Room.query.filter_by(cinema_id=cinema.id).one()
            assert room.name == "Phòng 01"
            assert room.room_type == "2D"
            assert Seat.query.filter_by(room_id=room.id).count() == 80

    public_cinemas = client.get("/api/cinemas").get_json()
    imported = {cinema["name"]: cinema for cinema in public_cinemas}
    assert imported["CGV Vincom Bà Triệu"]["region"] == "Miền Bắc"
    assert imported["CGV Vincom Đà Nẵng"]["region"] == "Miền Trung"
    assert imported["CGV Kim Cương Quy Nhơn"]["province"] == "Gia Lai"
    assert imported["CGV Aeon Bình Dương"]["province"] == "Thành phố Hồ Chí Minh"

    repeated = client.post("/api/admin/cinemas/import", headers=auth("admin"), data={
        "file": (BytesIO(file_content.getvalue()), "DanhSachRapChieu.xlsx"),
    })
    assert repeated.status_code == 200
    assert repeated.get_json()["imported_count"] == 0
    assert repeated.get_json()["skipped_count"] == 4


def test_import_backfills_chain_for_previously_imported_cinema(client, auth, app):
    with app.app_context():
        legacy = Cinema(
            name="Beta Bắc Giang",
            city="Bắc Ninh",
            address="Địa chỉ Beta",
        )
        db.session.add(legacy)
        db.session.commit()

    csv_data = (
        "Hệ thống rạp,Tên cụm rạp,Khu vực / Tỉnh thành,Địa chỉ / Chi tiết\n"
        "Beta Cinemas,Beta Bắc Giang,Miền Bắc (Bắc Giang),Địa chỉ Beta\n"
    ).encode("utf-8-sig")
    response = client.post("/api/admin/cinemas/import", headers=auth("admin"), data={
        "file": (BytesIO(csv_data), "cinemas.csv"),
    })
    assert response.status_code == 200
    assert response.get_json()["imported_count"] == 0
    assert response.get_json()["skipped_count"] == 1
    with app.app_context():
        cinema = Cinema.query.filter_by(name="Beta Bắc Giang").one()
        assert cinema.chain == "Beta Cinemas"
        assert Room.query.filter_by(cinema_id=cinema.id).count() == 0
