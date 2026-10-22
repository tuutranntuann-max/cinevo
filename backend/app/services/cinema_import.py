"""Read and validate cinema and room rows from administrator uploads."""
import csv
import re
from io import BytesIO, StringIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ..data.vietnam_provinces import normalize_province, region_for_province

MAX_CINEMA_IMPORT_ROWS = 500
MAX_XLSX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
HEADERS = {
    "tên rạp": "name",
    "tên cụm rạp": "name",
    "cinema": "name",
    "cinema name": "name",
    "name": "name",
    "hệ thống rạp": "chain",
    "chuỗi rạp": "chain",
    "rạp": "chain",
    "miền": "region",
    "region": "region",
    "khu vực / tỉnh thành": "city",
    "khu vực/tỉnh thành": "city",
    "khu vực tỉnh thành": "city",
    "tỉnh/thành phố": "city",
    "tỉnh/thành": "city",
    "tỉnh thành": "city",
    "tỉnh thành phố": "city",
    "tỉnh": "city",
    "thành phố": "city",
    "province": "city",
    "city": "city",
    "địa chỉ": "address",
    "địa chỉ / chi tiết": "address",
    "địa chỉ/chi tiết": "address",
    "address": "address",
    "sđt": "phone",
    "điện thoại": "phone",
    "phone": "phone",
    "tên phòng": "room_name",
    "phòng chiếu": "room_name",
    "room name": "room_name",
    "loại phòng": "room_type",
    "room type": "room_type",
    "số hàng": "total_rows",
    "rows": "total_rows",
    "số cột": "total_cols",
    "columns": "total_cols",
}


def _read_rows(content, extension):
    if extension == "csv":
        text = content.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = []
        for row in csv.reader(StringIO(text), dialect):
            rows.append(row)
            if len(rows) > MAX_CINEMA_IMPORT_ROWS + 1:
                raise ValueError(f"Mỗi lần chỉ nhập tối đa {MAX_CINEMA_IMPORT_ROWS} phòng.")
        return rows

    with ZipFile(BytesIO(content)) as archive:
        files = archive.infolist()
        if len(files) > 1000 or sum(item.file_size for item in files) > MAX_XLSX_UNCOMPRESSED_BYTES:
            raise ValueError("Nội dung giải nén của file Excel vượt quá giới hạn cho phép.")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        if worksheet.max_row and worksheet.max_row > MAX_CINEMA_IMPORT_ROWS + 1:
            raise ValueError(f"Mỗi lần chỉ nhập tối đa {MAX_CINEMA_IMPORT_ROWS} phòng.")
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _text(value):
    return "" if value is None else str(value).strip()


def _number(value, field, row_number, default, maximum):
    if value in (None, ""):
        return default
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"Dòng {row_number}: {field} phải là số nguyên.")
    if not number.is_integer():
        raise ValueError(f"Dòng {row_number}: {field} phải là số nguyên.")
    result = int(number)
    if not 1 <= result <= maximum:
        raise ValueError(f"Dòng {row_number}: {field} phải từ 1 đến {maximum}.")
    return result


def _location(value):
    text = _text(value)
    match = re.fullmatch(r"(miền\s+(?:bắc|trung|nam))\s*(?:\((.+)\))?", text, re.IGNORECASE)
    if match:
        region = " ".join(match.group(1).casefold().split()).title()
        province = match.group(2).strip() if match.group(2) else ""
        return region, normalize_province(province)
    return "", normalize_province(text)


def parse_cinema_import(content, extension):
    try:
        rows = _read_rows(content, extension)
    except (UnicodeDecodeError, csv.Error) as exc:
        raise ValueError("Không đọc được file. Hãy dùng CSV UTF-8 hoặc Excel .xlsx hợp lệ.") from exc
    except (BadZipFile, InvalidFileException, OSError, KeyError, IndexError) as exc:
        raise ValueError("Không đọc được file Excel. Hãy tải mẫu và thử lại.") from exc

    if not rows:
        raise ValueError("File không có dữ liệu.")
    headers = {}
    for index, cell in enumerate(rows[0]):
        key = HEADERS.get(_text(cell).casefold())
        if key:
            if key in headers:
                raise ValueError(f"Cột “{key}” bị lặp trong dòng tiêu đề.")
            headers[key] = index
    for required in ("name", "city"):
        if required not in headers:
            labels = {
                "name": "Tên rạp (hỗ trợ “Tên cụm rạp”)",
                "city": "Tỉnh/thành (hỗ trợ “Tỉnh/thành phố” hoặc “Khu vực / Tỉnh thành”)",
            }
            raise ValueError(f"Thiếu cột bắt buộc “{labels[required]}”.")

    cinemas = {}
    errors = []
    row_count = 0
    room_column_present = "room_name" in headers
    for row_number, row in enumerate(rows[1:], start=2):
        raw = {key: row[index] if index < len(row) else None for key, index in headers.items()}
        values = {key: _text(value) for key, value in raw.items()}
        if not any(values.values()):
            continue
        row_count += 1
        if row_count > MAX_CINEMA_IMPORT_ROWS:
            raise ValueError(f"Mỗi lần chỉ nhập tối đa {MAX_CINEMA_IMPORT_ROWS} phòng.")

        name = values.get("name", "")
        chain = values.get("chain", "") or None
        supplied_region, city = _location(values.get("city", ""))
        separate_region = values.get("region", "")
        if separate_region:
            separate_region = " ".join(separate_region.casefold().split()).title()
        supplied_region = separate_region or supplied_region
        address = values.get("address", "") or None
        phone = values.get("phone", "")
        room_name = values.get("room_name", "") or "Phòng 01"
        region = region_for_province(city) if city else None
        if not name or len(name) > 150:
            errors.append(f"Dòng {row_number}: thiếu tên rạp hoặc tên vượt quá 150 ký tự.")
            continue
        if chain and len(chain) > 100:
            errors.append(f"Dòng {row_number}: tên hệ thống rạp vượt quá 100 ký tự.")
            continue
        if not city:
            errors.append(f"Dòng {row_number}: tỉnh/thành phố không thuộc danh mục hiện hành 34 tỉnh/thành.")
            continue
        if address and len(address) > 250:
            errors.append(f"Dòng {row_number}: địa chỉ vượt quá 250 ký tự.")
            continue
        if len(phone) > 20:
            errors.append(f"Dòng {row_number}: số điện thoại vượt quá 20 ký tự.")
            continue
        if not room_name or len(room_name) > 50:
            errors.append(f"Dòng {row_number}: thiếu tên phòng hoặc tên vượt quá 50 ký tự.")
            continue

        room_type = values.get("room_type", "").upper() or "2D"
        if room_type not in {"2D", "3D", "IMAX"}:
            errors.append(f"Dòng {row_number}: loại phòng phải là 2D, 3D hoặc IMAX.")
            continue
        try:
            total_rows = _number(raw.get("total_rows"), "số hàng", row_number, 8, 26)
            total_cols = _number(raw.get("total_cols"), "số cột", row_number, 10, 30)
        except ValueError as exc:
            errors.append(str(exc))
            continue
        region_corrected = bool(
            supplied_region and supplied_region.casefold() != (region or "").casefold()
        )

        cinema_key = (
            " ".join(chain.casefold().split()) if chain else "",
            " ".join(name.casefold().split()),
            city,
        )
        if not room_column_present and cinema_key in cinemas:
            previous = cinemas[cinema_key]
            if previous["address"] == address and previous["phone"] == (phone or None):
                previous["region_corrected"] = previous["region_corrected"] or region_corrected
                continue
            errors.append(f"Dòng {row_number}: thông tin địa chỉ không đồng nhất với rạp trùng tên.")
            continue
        cinema = cinemas.setdefault(cinema_key, {
            "chain": chain,
            "name": name,
            "city": city,
            "region": region,
            "region_corrected": region_corrected,
            "address": address,
            "phone": phone or None,
            "rooms": [],
            "room_names": set(),
            "row_number": row_number,
        })
        if cinema["address"] != address or cinema["phone"] != (phone or None):
            errors.append(f"Dòng {row_number}: địa chỉ hoặc điện thoại không đồng nhất với các dòng cùng rạp.")
            continue
        cinema["region_corrected"] = cinema["region_corrected"] or region_corrected
        room_key = room_name.casefold()
        if room_key in cinema["room_names"]:
            errors.append(f"Dòng {row_number}: tên phòng “{room_name}” bị lặp trong cùng một rạp.")
            continue
        cinema["room_names"].add(room_key)
        cinema["rooms"].append({
            "name": room_name,
            "room_type": room_type,
            "total_rows": total_rows,
            "total_cols": total_cols,
        })

    if errors:
        summary = " ".join(errors[:10])
        remaining = len(errors) - 10
        if remaining:
            summary += f" Và còn {remaining} dòng có lỗi."
        raise ValueError(summary)
    if not cinemas:
        raise ValueError("File không có rạp/phòng chiếu hợp lệ để nhập.")
    return list(cinemas.values())
